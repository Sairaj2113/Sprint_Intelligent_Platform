from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document as DocxDocument

from app.services.document_extraction_service import (
    NoExtractableTextError,
    extract_document,
)


class DocumentExtractionServiceTests(unittest.TestCase):
    def _pdf_path(self, root: Path) -> tuple[Path, str]:
        relative_path = "RD/document-id/requirements.pdf"
        path = root / relative_path
        path.parent.mkdir(parents=True)
        path.write_bytes(b"placeholder PDF for mocked reader")
        return path, relative_path

    def test_pdf_text_is_extracted_in_page_order_and_blank_pages_are_ignored(self) -> None:
        with tempfile.TemporaryDirectory() as temp_directory:
            root = Path(temp_directory)
            _, relative_path = self._pdf_path(root)
            reader = SimpleNamespace(
                pages=[
                    SimpleNamespace(extract_text=lambda: " First page\r\n\r\n\r\ntext "),
                    SimpleNamespace(extract_text=lambda: "   "),
                    SimpleNamespace(extract_text=lambda: "Third page"),
                ]
            )
            with patch(
                "app.services.document_extraction_service.PdfReader", return_value=reader
            ):
                blocks = extract_document(
                    storage_root=root,
                    storage_path=relative_path,
                    original_filename="requirements.pdf",
                )

        self.assertEqual([block.page_number for block in blocks], [1, 3])
        self.assertEqual([block.text for block in blocks], ["First page\n\ntext", "Third page"])
        self.assertEqual(blocks[0].metadata, {"source_type": "pdf", "page_number": 1})

    def test_textless_pdf_raises_controlled_error(self) -> None:
        with tempfile.TemporaryDirectory() as temp_directory:
            root = Path(temp_directory)
            _, relative_path = self._pdf_path(root)
            reader = SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda: "")])
            with patch(
                "app.services.document_extraction_service.PdfReader", return_value=reader
            ), self.assertRaises(NoExtractableTextError):
                extract_document(
                    storage_root=root,
                    storage_path=relative_path,
                    original_filename="requirements.pdf",
                )

    def test_docx_paragraphs_preserve_order_headings_and_skip_blanks(self) -> None:
        with tempfile.TemporaryDirectory() as temp_directory:
            root = Path(temp_directory)
            relative_path = "RD/document-id/architecture.docx"
            path = root / relative_path
            path.parent.mkdir(parents=True)
            document = DocxDocument()
            document.add_heading("Architecture", level=1)
            document.add_paragraph("FastAPI serves RD-15.")
            document.add_paragraph("   ")
            document.add_heading("Security", level=2)
            document.add_paragraph("AES-256-GCM protects stored data.")
            document.save(path)

            blocks = extract_document(
                storage_root=root,
                storage_path=relative_path,
                original_filename="architecture.docx",
            )

        self.assertEqual(
            [block.text for block in blocks],
            [
                "Architecture",
                "FastAPI serves RD-15.",
                "Security",
                "AES-256-GCM protects stored data.",
            ],
        )
        self.assertEqual(
            [block.section_title for block in blocks],
            ["Architecture", "Architecture", "Security", "Security"],
        )
        self.assertTrue(all(block.page_number is None for block in blocks))
        self.assertEqual(blocks[-1].metadata["source_type"], "docx")


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
